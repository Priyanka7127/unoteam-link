"""
report_builder.py — turns crawl results into the same audit workbook
structure used throughout this engagement: Summary, Treatments, Blogs,
Case Studies tabs, no Category column, a Score column showing the
similarity strength behind each recommendation, and cross-type Scope
recommendations (with a bolded suggested anchor text) for every orphan page.
"""
import re
from io import BytesIO
from urllib.parse import urlparse

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.text import InlineFont
from openpyxl.cell.rich_text import TextBlock, CellRichText

import crawler_core as cc

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14, color="1F4E78")
SUBTITLE_FONT = Font(name=FONT_NAME, italic=True, size=9, color="595959")
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
WRAP_C = Alignment(wrap_text=True, vertical="top", horizontal="center")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RED_FILL = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
GREEN_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
LIGHT_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
BOLD_INLINE = InlineFont(rFont=FONT_NAME, sz=9.5, b=True)


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP_C
        cell.border = BORDER


def _style_row(ws, row, ncols, fill=None):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT_NAME, size=9.5)
        cell.alignment = WRAP
        cell.border = BORDER
        if fill:
            cell.fill = fill


def _set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title_block(ws, title, subtitle, ncols, sub_height=40):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = sub_height


def _clean_anchor_text(title: str) -> str:
    """Turn a raw <title> into anchor-ready text: drop the trailing
    '| Brand Name' / '- Brand Name' boilerplate that page titles carry,
    so the suggested anchor reads naturally instead of a whole SEO title."""
    if not title:
        return title
    for sep in (" | ", " — ", " – ", " - "):
        if sep in title:
            head = title.split(sep)[0].strip()
            if len(head) >= 8:  # avoid chopping down to something too short/generic
                return head
    return title.strip()


_LOCATION_SUFFIX_RE = re.compile(
    r"\s+(?:in|for|near)\s+[A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+){0,2}$"
)


def _shorten_phrase(title: str) -> str:
    """Drop a trailing ' in <City>' / ' for <Place>' geo-suffix so the core
    topic phrase stays short — SEO anchor text should be a small, natural
    phrase, not a whole keyword-stuffed page title."""
    cleaned = _clean_anchor_text(title)
    shortened = _LOCATION_SUFFIX_RE.sub("", cleaned).strip()
    return shortened or cleaned


def _anchor_variants(title: str):
    """A small rotation of natural, short anchor phrasings for the same
    target page — using one identical exact-match anchor on every single
    recommended link looks manipulative to Google; real internal-linking
    profiles vary anchor text (exact/partial-match, natural mentions)."""
    base = _shorten_phrase(title)
    return [base, base.lower(), f"more on {base.lower()}"]


def _format_scope(recs, target_title, needed=None):
    """Build the Scope cell as rich text: plain-text recommendation details
    with a short, naturally varied suggested anchor text bolded inline per URL —
    not the same exact-match full title repeated on every recommendation.
    If fewer relevant candidates exist than are needed to close the gap,
    say so explicitly rather than silently padding with weaker matches."""
    if not recs:
        note = "No sufficiently relevant page found automatically."
        if needed:
            note += f" This page still needs {needed} more link(s) to reach its target — review manually rather than adding an unrelated link just to hit the count."
        else:
            note += " Review manually."
        return note
    variants = _anchor_variants(target_title)
    header = "Consider adding links from:"
    if needed is not None and len(recs) < needed:
        header += (f" (only {len(recs)} sufficiently relevant page(s) found — fewer than the "
                   f"{needed} needed to close the gap; don't pad with unrelated pages)")
    pieces = [header]
    for i, (score, url, title, category) in enumerate(recs):
        anchor = variants[i % len(variants)]
        pieces.append(f"\n[{category}] \"{title}\" ({url}) — score {score:.2f} — anchor text: ")
        pieces.append(TextBlock(BOLD_INLINE, anchor))
    return CellRichText(pieces)


def find_broken_links(pages: dict, inbound: dict):
    """Pages that were linked to but returned a non-200 status or failed
    outright — real broken-link findings, not content to list as pages."""
    broken = []
    for url, data in pages.items():
        if data.get("error") or (data.get("status") is not None and data["status"] != 200):
            sources = inbound.get(url, [])
            broken.append({
                "url": url,
                "status": data.get("status"),
                "error": data.get("error"),
                "linked_from": sources,
            })
    return broken


def _make_category_sheet(wb, sheet_name, title, subtitle, pages, inbound, category_keys, min_links):
    rows = {u: d for u, d in pages.items()
            if d["category"] in category_keys and d.get("status") == 200}
    ws = wb.create_sheet(sheet_name)
    headers = ["#", "Page Title", "URL", "Current Inbound Links",
               "Linking Page(s) & Anchor Text", f"Links Needed (target: {min_links})", "Score",
               "Scope: Recommended Links to Add (relevant only — never padded to hit a count)"]
    _title_block(ws, title, subtitle, len(headers))
    hr = 3
    for i, h in enumerate(headers, start=1):
        ws.cell(row=hr, column=i, value=h)
    _style_header(ws, hr, len(headers))
    ws.freeze_panes = f"A{hr + 1}"

    r = hr + 1
    below_target = 0
    for idx, (url, data) in enumerate(sorted(rows.items(), key=lambda kv: kv[1]["title"]), start=1):
        links_in = inbound.get(url, [])
        count = len(links_in)
        evidence = "; ".join(f"{s['source']} → \"{s['anchor_text']}\"" for s in links_in) or "None found in this crawl."
        needed = max(0, min_links - count)
        score = ""
        if needed > 0:
            below_target += 1
            existing_sources = {s["source"] for s in links_in}
            recs = cc.recommend_link_sources(url, pages, top_n=needed, exclude_urls=existing_sources)
            score = round(recs[0][0], 2) if recs else 0
            scope = _format_scope(recs, data["title"], needed=needed)
        else:
            scope = f"Already has {count} inbound link(s), meeting the {min_links}-link target for this page type."
        row_vals = [idx, data["title"], url, count, evidence, needed, score, scope]
        for c, v in enumerate(row_vals, start=1):
            ws.cell(row=r, column=c, value=v)
        fill = RED_FILL if count == 0 else (YELLOW_FILL if needed > 0 else GREEN_FILL)
        _style_row(ws, r, len(headers), fill=fill)
        ws.cell(row=r, column=1).alignment = WRAP_C
        ws.cell(row=r, column=4).alignment = WRAP_C
        ws.cell(row=r, column=6).alignment = WRAP_C
        ws.cell(row=r, column=7).alignment = WRAP_C
        r += 1

    _set_widths(ws, [4, 30, 34, 12, 46, 12, 10, 56])
    for row in ws.iter_rows(min_row=hr + 1, max_row=r - 1):
        ws.row_dimensions[row[0].row].height = 60
    if r > hr + 1:
        ws.auto_filter.ref = f"A{hr}:H{r - 1}"
    return len(rows), below_target


def _make_homepage_link_sheet(wb, pages, inbound, site_url, selected):
    """Every page on the site should carry at least one contextual (non-nav)
    link back to the homepage — it's the page most worth reinforcing. This
    sheet flags every audited page that currently doesn't."""
    home_url = next((u for u, d in pages.items() if d["category"] == "home" and d.get("status") == 200), None)
    if home_url is None:
        return None, 0, 0
    home_anchor = "Home"
    linked_sources = {s["source"] for s in inbound.get(home_url, [])}

    wanted_cats = {cat for cat, on in
                   (("content", selected.get("content")), ("blog", selected.get("blog")),
                    ("case_study", selected.get("case_study"))) if on}
    rows = {u: d for u, d in pages.items()
            if d["category"] in wanted_cats and d.get("status") == 200 and u != home_url}

    ws = wb.create_sheet("Homepage Linking")
    headers = ["#", "Page Title", "URL", "Links to Home?", "Suggested Anchor Text"]
    _title_block(ws, "Homepage Internal Linking Check",
                  f"Crawled from {site_url}. Every page on a site should carry at least one contextual "
                  "body link back to the homepage — it consolidates link equity on the page you most "
                  "want to rank and helps users/crawlers navigate back. This checks contextual "
                  "(non-nav/footer) links only, same methodology as the other tabs.",
                  len(headers))
    hr = 3
    for i, h in enumerate(headers, start=1):
        ws.cell(row=hr, column=i, value=h)
    _style_header(ws, hr, len(headers))
    ws.freeze_panes = f"A{hr + 1}"

    r = hr + 1
    missing = 0
    for idx, (url, data) in enumerate(sorted(rows.items(), key=lambda kv: kv[1]["title"]), start=1):
        has_link = url in linked_sources
        if not has_link:
            missing += 1
        row_vals = [idx, data["title"], url, "Yes" if has_link else "NO",
                    "" if has_link else home_anchor]
        for c, v in enumerate(row_vals, start=1):
            ws.cell(row=r, column=c, value=v)
        if not has_link:
            ws.cell(row=r, column=5).font = Font(name=FONT_NAME, size=9.5, bold=True)
        fill = GREEN_FILL if has_link else RED_FILL
        _style_row(ws, r, len(headers), fill=fill)
        ws.cell(row=r, column=1).alignment = WRAP_C
        ws.cell(row=r, column=4).alignment = WRAP_C
        r += 1

    _set_widths(ws, [4, 34, 40, 14, 24])
    for row in ws.iter_rows(min_row=hr + 1, max_row=r - 1):
        ws.row_dimensions[row[0].row].height = 30
    if r > hr + 1:
        ws.auto_filter.ref = f"A{hr}:E{r - 1}"
    return home_url, len(rows), missing


def build_workbook(pages: dict, selected: dict, site_url: str) -> BytesIO:
    """selected: {'content': bool, 'blog': bool, 'case_study': bool}"""
    inbound = cc.build_inbound_index(pages)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    stats = {}
    LINK_TARGETS = {"content": 10, "blog": 4, "case_study": 4}

    if selected.get("content"):
        n, below = _make_category_sheet(
            wb, "Treatments-Services", "Treatments / Service Pages — Internal Linking Audit",
            f"Crawled from {site_url}. 'Current Inbound Links' counts contextual body links only "
            "(the site's repeating nav menu and footer are excluded, since those appear on every "
            "page and don't reflect real topical relevance). Target: {n} inbound links per service "
            "page. Any page below that gets ONLY genuinely relevant recommendations to close the gap "
            "— never padded with unrelated pages just to hit the count. Service pages are "
            "intentionally nudged toward linking from relevant blog posts or case studies rather than "
            "other service pages (a service page doesn't need to link to another service page just "
            "because it's the same type) — a same-type candidate is still used when it's the clearly "
            "stronger match.".format(n=LINK_TARGETS["content"]),
            pages, inbound, {"content"}, LINK_TARGETS["content"])
        stats["Treatments/Services"] = (n, below)

    if selected.get("blog"):
        n, below = _make_category_sheet(
            wb, "Blogs", "Blogs — Cross-Linking Audit",
            f"Crawled from {site_url}. Same methodology as the Treatments tab: contextual links only, "
            f"target of {LINK_TARGETS['blog']} inbound links per post, cross-type recommendations "
            "limited to genuinely relevant pages only.",
            pages, inbound, {"blog"}, LINK_TARGETS["blog"])
        stats["Blogs"] = (n, below)

    if selected.get("case_study"):
        n, below = _make_category_sheet(
            wb, "Case Studies", "Case Studies — Cross-Linking Audit",
            f"Crawled from {site_url}. Same methodology as the other tabs, target of "
            f"{LINK_TARGETS['case_study']} inbound links per case study.",
            pages, inbound, {"case_study"}, LINK_TARGETS["case_study"])
        stats["Case Studies"] = (n, below)

    # --- Homepage Linking sheet ---
    home_url, home_checked, home_missing = _make_homepage_link_sheet(wb, pages, inbound, site_url, selected)

    # --- Broken Links sheet ---
    broken = find_broken_links(pages, inbound)
    if broken:
        ws_b = wb.create_sheet("Broken Links")
        headers_b = ["#", "Broken URL", "Status / Error", "Linked From (page → anchor text)"]
        _title_block(ws_b, "Broken Internal Links Found",
                      f"Crawled from {site_url}. These are links this crawl followed that did not "
                      "return a working page (404, error, or non-HTML response) — each one is real "
                      "link equity currently going nowhere.",
                      len(headers_b))
        hr_b = 3
        for i, h in enumerate(headers_b, start=1):
            ws_b.cell(row=hr_b, column=i, value=h)
        _style_header(ws_b, hr_b, len(headers_b))
        ws_b.freeze_panes = f"A{hr_b + 1}"
        r = hr_b + 1
        for idx, item in enumerate(broken, start=1):
            status_txt = item["error"] or f"HTTP {item['status']}"
            sources_txt = "; ".join(f"{s['source']} → \"{s['anchor_text']}\"" for s in item["linked_from"]) or "(only reached via nav/footer or another broken page)"
            for c, v in enumerate([idx, item["url"], status_txt, sources_txt], start=1):
                ws_b.cell(row=r, column=c, value=v)
            _style_row(ws_b, r, len(headers_b), fill=RED_FILL)
            ws_b.cell(row=r, column=1).alignment = WRAP_C
            r += 1
        _set_widths(ws_b, [4, 40, 20, 60])
        for row in ws_b.iter_rows(min_row=hr_b + 1, max_row=r - 1):
            ws_b.row_dimensions[row[0].row].height = 40
        ws_b.auto_filter.ref = f"A{hr_b}:D{r - 1}"

    # --- Summary sheet (inserted first) ---
    ws = wb.create_sheet("Summary", 0)
    ws.sheet_view.showGridLines = False
    _set_widths(ws, [3, 46, 16, 46, 16])
    row = 1
    ws.merge_cells(f"A{row}:E{row}")
    ws.cell(row=row, column=1, value="Internal Linking Audit — Summary").font = TITLE_FONT
    row += 1
    ws.merge_cells(f"A{row}:E{row}")
    ws.cell(row=row, column=1, value=f"Site: {site_url}").font = SUBTITLE_FONT
    row += 2

    def section(title, r):
        ws.merge_cells(f"A{r}:E{r}")
        c = ws.cell(row=r, column=1, value=title)
        c.font = Font(name=FONT_NAME, bold=True, size=12, color="1F4E78")
        c.fill = LIGHT_FILL
        for col in range(1, 6):
            ws.cell(row=r, column=col).fill = LIGHT_FILL
            ws.cell(row=r, column=col).border = BORDER
        return r + 1

    def kv(r, label, value):
        ws.cell(row=r, column=2, value=label).font = Font(name=FONT_NAME, size=10)
        ws.cell(row=r, column=2).alignment = WRAP
        ws.cell(row=r, column=4, value=value).font = Font(name=FONT_NAME, size=10, bold=True)
        ws.cell(row=r, column=4).alignment = WRAP_C
        for col in [2, 3, 4, 5]:
            ws.cell(row=r, column=col).border = BORDER
        return r + 1

    row = section("Site Inventory (this crawl)", row)
    row = kv(row, "Total pages crawled", str(len(pages)))
    for label, (n, below) in stats.items():
        row = kv(row, f"{label} pages found", str(n))
        row = kv(row, f"{label} pages below their inbound-link target", str(below))
    row = kv(row, "Broken internal links found", str(len(broken)))
    if home_url is not None:
        row = kv(row, "Pages checked for a homepage link", str(home_checked))
        row = kv(row, "Pages missing a contextual link to homepage", str(home_missing))
    row += 1

    row = section("Methodology & Honest Limitations", row)
    ws.merge_cells(f"A{row}:E{row + 9}")
    note = (
        "This report was generated by an automated crawler running from your own computer, not a "
        "browser page — it made real HTTP requests to every page it found on the site, up to the "
        "page limit you set, following links breadth-first from the homepage.\n\n"
        "'Inbound links' counts only contextual links found in the page body; the site's repeating "
        "navigation menu and footer are excluded on purpose. If the crawl hit its page limit before "
        "reaching every corner of the site, some pages below their link target may actually have "
        "inbound links from pages that weren't reached this time. Increase the page limit and re-run "
        "for full confidence on a large site.\n\n"
        "Each page type has an inbound-link target: 10 for Treatments/Service pages, 4 for Blogs and "
        "Case Studies. Any page below its target gets recommendations to close the gap — but ONLY "
        "pages that clear a minimum topical-relevance score are ever recommended; if fewer relevant "
        "pages exist than are needed, the report says so explicitly rather than padding the list with "
        "unrelated pages. Service pages are intentionally nudged toward linking from a relevant blog "
        "post or case study rather than another service page, since two service pages being the same "
        "type doesn't make them a natural link pair — a same-type page is still recommended when it's "
        "clearly the strongest match.\n\n"
        "Scope recommendations are computed by simple keyword-overlap similarity between page titles "
        "and text, not by a human reading each page — they are candidates worth reviewing, not "
        "guaranteed-correct editorial judgments. The 'Score' column is the best recommendation's "
        "similarity value (0–1) — treat low scores with more skepticism than high ones. The suggested "
        "anchor text in bold within the Scope column is a short, cleaned phrase from the target page's "
        "own title (SEO-suffix like '| Brand Name' and city/location suffixes stripped), varied across "
        "recommendations rather than repeating one exact phrase on every link.\n\n"
        "The Homepage Linking tab checks every audited page for at least one contextual (non-nav) link "
        "back to the homepage — this consolidates link equity on your most important page. Any page "
        "marked NO should get a link added, using the suggested anchor text shown."
    )
    cell = ws.cell(row=row, column=1, value=note)
    cell.font = Font(name=FONT_NAME, size=9.5, italic=True)
    cell.alignment = WRAP
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
